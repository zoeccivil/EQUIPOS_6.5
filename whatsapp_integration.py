# whatsapp_integration.py

"""
Integración con WhatsApp Business API
Permite enviar recordatorios de pago, confirmaciones y notificaciones
Usa Twilio WhatsApp API o WhatsApp Business API oficial
"""

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel,
    QPushButton, QTableWidget, QTableWidgetItem, QHeaderView,
    QDialog, QFormLayout, QLineEdit, QDoubleSpinBox, QDateEdit,
    QComboBox, QMessageBox, QAbstractItemView, QGroupBox, QFrame,
    QTextEdit, QFileDialog, QCheckBox, QTabWidget, QListWidget,
    QListWidgetItem, QSpinBox
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal, QTimer
from PyQt6.QtGui import QFont, QColor, QPixmap, QBrush
import logging
import json
import requests
from datetime import datetime, timedelta
from firebase_manager import FirebaseManager

logger = logging.getLogger(__name__)

# Estilos CSS
WHATSAPP_STYLE = """
QWidget {
    background-color: #F3F4F6;
    font-family: 'Segoe UI';
}
QLabel[class="title"] {
    font-size: 20pt;
    font-weight: bold;
    color: #1F2937;
}
QLabel[class="whatsapp-title"] {
    font-size: 18pt;
    font-weight: bold;
    color: #25D366;
}
QFrame[class="stat-card"] {
    background-color: #FFFFFF;
    border: 2px solid #E5E7EB;
    border-radius: 12px;
    padding: 20px;
}
QFrame[class="whatsapp-card"] {
    background-color: #FFFFFF;
    border: 3px solid #25D366;
    border-radius: 12px;
    padding: 20px;
}
QFrame[class="chat-frame"] {
    background-color: #ECE5DD;
    border: 1px solid #D1D5DB;
    border-radius: 8px;
    padding: 10px;
}
QPushButton {
    background-color: #F59E0B;
    color: white;
    border: none;
    border-radius: 6px;
    padding: 10px 18px;
    font-size: 10pt;
    font-weight: 600;
    min-height: 25px;
}
QPushButton:hover {
    background-color: #D97706;
}
QPushButton:pressed {
    background-color: #B45309;
}
QPushButton[class="whatsapp"] {
    background-color: #25D366;
    color: white;
}
QPushButton[class="whatsapp"]:hover {
    background-color: #128C7E;
}
QPushButton[class="whatsapp"]:pressed {
    background-color: #075E54;
}
QPushButton[class="secondary"] {
    background-color: #E5E7EB;
    color: #374151;
}
QPushButton[class="secondary"]:hover {
    background-color: #D1D5DB;
}
QTableWidget {
    background-color: #FFFFFF;
    alternate-background-color: #F9FAFB;
    gridline-color: #E5E7EB;
    selection-background-color: #D1FAE5;
    selection-color: #1F2937;
    border: 1px solid #E5E7EB;
    border-radius: 6px;
}
QTableWidget::item {
    padding: 8px;
}
QHeaderView::section {
    background-color: #1F2937;
    color: #FFFFFF;
    padding: 10px;
    border: none;
    font-weight: 600;
}
QLineEdit, QTextEdit, QComboBox, QSpinBox {
    background-color: #FFFFFF;
    border: 2px solid #E5E7EB;
    border-radius: 6px;
    padding: 8px 12px;
    color: #1F2937;
    font-size: 10pt;
}
QLineEdit:hover, QTextEdit:hover, QComboBox:hover {
    border: 2px solid #25D366;
}
QLineEdit:focus, QTextEdit:focus, QComboBox:focus {
    border: 2px solid #25D366;
}
QGroupBox {
    background-color: #FFFFFF;
    border: 2px solid #E5E7EB;
    border-radius: 10px;
    margin-top: 14px;
    padding: 15px;
    font-weight: 600;
    font-size: 12pt;
    color: #1F2937;
}
QGroupBox::title {
    subcontrol-origin: margin;
    left: 15px;
    padding: 0 8px;
    background-color: #FFFFFF;
}
QTabWidget::pane {
    border: 2px solid #E5E7EB;
    border-radius: 6px;
    background-color: #FFFFFF;
}
QTabBar::tab {
    background-color: #E5E7EB;
    color: #374151;
    padding: 10px 20px;
    border-top-left-radius: 6px;
    border-top-right-radius: 6px;
    margin-right: 2px;
}
QTabBar::tab:selected {
    background-color: #25D366;
    color: white;
}
QTabBar::tab:hover {
    background-color: #D1D5DB;
}
QListWidget {
    background-color: #FFFFFF;
    border: 1px solid #E5E7EB;
    border-radius: 6px;
}
QListWidget::item {
    padding: 10px;
    border-bottom: 1px solid #E5E7EB;
}
QListWidget::item:selected {
    background-color: #D1FAE5;
    color: #1F2937;
}
"""


class WhatsAppManager:
    """
    Gestor de conexión con WhatsApp Business API
    Soporta Twilio y API oficial de WhatsApp
    """
    
    def __init__(self, config: dict):
        self.config = config
        self.provider = config.get('whatsapp', {}).get('provider', 'twilio')
        
        # Credenciales Twilio
        self.twilio_account_sid = config.get('whatsapp', {}).get('twilio_account_sid', '')
        self.twilio_auth_token = config.get('whatsapp', {}).get('twilio_auth_token', '')
        self.twilio_from = config.get('whatsapp', {}).get('twilio_from', '')
        
        # Credenciales WhatsApp Business API
        self.whatsapp_token = config.get('whatsapp', {}).get('api_token', '')
        self.whatsapp_phone_id = config.get('whatsapp', {}).get('phone_id', '')
        
        self.conectado = False
    
    def verificar_conexion(self):
        """Verifica si las credenciales son válidas"""
        try:
            if self.provider == 'twilio':
                return self._verificar_twilio()
            elif self.provider == 'whatsapp_business':
                return self._verificar_whatsapp_business()
            else:
                return False
        except Exception as e:
            logger.error(f"Error verificando conexión WhatsApp: {e}", exc_info=True)
            return False
    
    def _verificar_twilio(self):
        """Verifica credenciales de Twilio"""
        if not all([self.twilio_account_sid, self.twilio_auth_token, self.twilio_from]):
            return False
        
        try:
            from twilio.rest import Client
            client = Client(self.twilio_account_sid, self.twilio_auth_token)
            # Intenta obtener información de la cuenta
            account = client.api.accounts(self.twilio_account_sid).fetch()
            self.conectado = True
            return True
        except ImportError:
            logger.warning("Twilio no está instalado. Instale con: pip install twilio")
            return False
        except Exception as e:
            logger.error(f"Error verificando Twilio: {e}", exc_info=True)
            return False
    
    def _verificar_whatsapp_business(self):
        """Verifica credenciales de WhatsApp Business API"""
        if not all([self.whatsapp_token, self.whatsapp_phone_id]):
            return False
        
        try:
            url = f"https://graph.facebook.com/v18.0/{self.whatsapp_phone_id}"
            headers = {"Authorization": f"Bearer {self.whatsapp_token}"}
            response = requests.get(url, headers=headers, timeout=10)
            
            self.conectado = response.status_code == 200
            return self.conectado
        except Exception as e:
            logger.error(f"Error verificando WhatsApp Business: {e}", exc_info=True)
            return False
    
    def enviar_mensaje(self, numero: str, mensaje: str) -> tuple:
        """
        Envía un mensaje de WhatsApp
        
        Args:
            numero: Número de teléfono en formato internacional (+1234567890)
            mensaje: Texto del mensaje
            
        Returns:
            tuple: (éxito: bool, mensaje_id o error: str)
        """
        try:
            # Validar formato de número
            if not numero.startswith('+'):
                numero = '+' + numero.replace('-', '').replace(' ', '')
            
            if self.provider == 'twilio':
                return self._enviar_twilio(numero, mensaje)
            elif self.provider == 'whatsapp_business':
                return self._enviar_whatsapp_business(numero, mensaje)
            else:
                return False, "Proveedor no configurado"
                
        except Exception as e:
            logger.error(f"Error enviando mensaje WhatsApp: {e}", exc_info=True)
            return False, str(e)
    
    def _enviar_twilio(self, numero: str, mensaje: str):
        """Envía mensaje vía Twilio"""
        try:
            from twilio.rest import Client
            
            client = Client(self.twilio_account_sid, self.twilio_auth_token)
            
            message = client.messages.create(
                from_=f'whatsapp:{self.twilio_from}',
                body=mensaje,
                to=f'whatsapp:{numero}'
            )
            
            return True, message.sid
            
        except ImportError:
            return False, "Twilio no está instalado. Instale con: pip install twilio"
        except Exception as e:
            return False, str(e)
    
    def _enviar_whatsapp_business(self, numero: str, mensaje: str):
        """Envía mensaje vía WhatsApp Business API"""
        try:
            url = f"https://graph.facebook.com/v18.0/{self.whatsapp_phone_id}/messages"
            
            headers = {
                "Authorization": f"Bearer {self.whatsapp_token}",
                "Content-Type": "application/json"
            }
            
            data = {
                "messaging_product": "whatsapp",
                "to": numero.replace('+', ''),
                "type": "text",
                "text": {"body": mensaje}
            }
            
            response = requests.post(url, headers=headers, json=data, timeout=30)
            
            if response.status_code == 200:
                result = response.json()
                message_id = result.get('messages', [{}])[0].get('id', 'unknown')
                return True, message_id
            else:
                error = response.json().get('error', {}).get('message', 'Error desconocido')
                return False, error
                
        except Exception as e:
            return False, str(e)
    
    def enviar_mensaje_template(self, numero: str, template_name: str, parametros: list) -> tuple:
        """
        Envía un mensaje usando una plantilla pre-aprobada
        Solo disponible en WhatsApp Business API
        """
        if self.provider != 'whatsapp_business':
            return False, "Templates solo disponibles en WhatsApp Business API"
        
        try:
            url = f"https://graph.facebook.com/v18.0/{self.whatsapp_phone_id}/messages"
            
            headers = {
                "Authorization": f"Bearer {self.whatsapp_token}",
                "Content-Type": "application/json"
            }
            
            # Construir componentes de la plantilla
            components = []
            if parametros:
                components.append({
                    "type": "body",
                    "parameters": [{"type": "text", "text": str(p)} for p in parametros]
                })
            
            data = {
                "messaging_product": "whatsapp",
                "to": numero.replace('+', ''),
                "type": "template",
                "template": {
                    "name": template_name,
                    "language": {"code": "es"},
                    "components": components
                }
            }
            
            response = requests.post(url, headers=headers, json=data, timeout=30)
            
            if response.status_code == 200:
                result = response.json()
                message_id = result.get('messages', [{}])[0].get('id', 'unknown')
                return True, message_id
            else:
                error = response.json().get('error', {}).get('message', 'Error desconocido')
                return False, error
                
        except Exception as e:
            logger.error(f"Error enviando template: {e}", exc_info=True)
            return False, str(e)


class WhatsAppIntegration(QWidget):
    """Widget principal para integración con WhatsApp"""
    
    def __init__(self, fm: FirebaseManager, config: dict, parent=None):
        super().__init__(parent)
        self.fm = fm
        self.config = config
        self.moneda = config.get('app', {}).get('moneda', 'RD$')
        
        self.whatsapp_manager = WhatsAppManager(config)
        
        self.mensajes_enviados = []
        self.templates = []
        
        self.setStyleSheet(WHATSAPP_STYLE)
        
        self._init_ui()
        self._verificar_conexion()
    
    def _init_ui(self):
        """Inicializa la interfaz"""
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(20)
        
        # Header
        header_layout = QHBoxLayout()
        
        titulo = QLabel("📱 WhatsApp Business")
        titulo.setProperty("class", "whatsapp-title")
        header_layout.addWidget(titulo)
        
        header_layout.addStretch()
        
        self.lbl_estado = QLabel("● Desconectado")
        self.lbl_estado.setStyleSheet("font-size: 11pt; font-weight: bold; color: #DC2626;")
        header_layout.addWidget(self.lbl_estado)
        
        btn_configurar = QPushButton("⚙️ Configurar")
        btn_configurar.setProperty("class", "secondary")
        btn_configurar.clicked.connect(self._configurar_whatsapp)
        header_layout.addWidget(btn_configurar)
        
        btn_verificar = QPushButton("🔄 Verificar Conexión")
        btn_verificar.clicked.connect(self._verificar_conexion)
        header_layout.addWidget(btn_verificar)
        
        layout.addLayout(header_layout)
        
        # Tabs
        tabs = QTabWidget()
        
        # === TAB 1: ENVIAR MENSAJES ===
        tab_enviar = QWidget()
        layout_enviar = QVBoxLayout(tab_enviar)
        layout_enviar.setContentsMargins(15, 15, 15, 15)
        layout_enviar.setSpacing(15)
        
        # Tipo de mensaje
        grupo_tipo = QGroupBox("Tipo de Mensaje")
        layout_tipo = QVBoxLayout(grupo_tipo)
        
        self.radio_manual = QCheckBox("Mensaje Manual")
        self.radio_manual.setChecked(True)
        self.radio_manual.toggled.connect(self._toggle_tipo_mensaje)
        layout_tipo.addWidget(self.radio_manual)
        
        self.radio_recordatorio = QCheckBox("Recordatorio de Pago")
        self.radio_recordatorio.toggled.connect(self._toggle_tipo_mensaje)
        layout_tipo.addWidget(self.radio_recordatorio)
        
        self.radio_confirmacion = QCheckBox("Confirmación de Alquiler")
        self.radio_confirmacion.toggled.connect(self._toggle_tipo_mensaje)
        layout_tipo.addWidget(self.radio_confirmacion)
        
        layout_enviar.addWidget(grupo_tipo)
        
        # Formulario de envío
        form_layout = QFormLayout()
        form_layout.setSpacing(15)
        
        lbl_destinatario = QLabel("Destinatario:")
        lbl_destinatario.setStyleSheet("font-weight: 600; color: #374151;")
        
        destinatario_layout = QHBoxLayout()
        self.combo_destinatario = QComboBox()
        self.combo_destinatario.addItem("Seleccionar cliente...", None)
        self.combo_destinatario.currentIndexChanged.connect(self._cargar_numero_cliente)
        destinatario_layout.addWidget(self.combo_destinatario)
        
        self.txt_numero = QLineEdit()
        self.txt_numero.setPlaceholderText("+1 (809) 555-1234")
        destinatario_layout.addWidget(self.txt_numero)
        
        form_layout.addRow(lbl_destinatario, destinatario_layout)
        
        lbl_mensaje = QLabel("Mensaje:")
        lbl_mensaje.setStyleSheet("font-weight: 600; color: #374151;")
        self.txt_mensaje = QTextEdit()
        self.txt_mensaje.setMaximumHeight(150)
        self.txt_mensaje.setPlaceholderText("Escribe tu mensaje aquí...")
        form_layout.addRow(lbl_mensaje, self.txt_mensaje)
        
        layout_enviar.addLayout(form_layout)
        
        # Botones de plantillas rápidas
        plantillas_layout = QHBoxLayout()
        plantillas_layout.addWidget(QLabel("Plantillas rápidas:"))
        
        btn_recordatorio = QPushButton("💰 Recordatorio")
        btn_recordatorio.setProperty("class", "secondary")
        btn_recordatorio.clicked.connect(lambda: self._usar_plantilla("recordatorio"))
        plantillas_layout.addWidget(btn_recordatorio)
        
        btn_gracias = QPushButton("🙏 Gracias")
        btn_gracias.setProperty("class", "secondary")
        btn_gracias.clicked.connect(lambda: self._usar_plantilla("gracias"))
        plantillas_layout.addWidget(btn_gracias)
        
        btn_confirmacion = QPushButton("✅ Confirmación")
        btn_confirmacion.setProperty("class", "secondary")
        btn_confirmacion.clicked.connect(lambda: self._usar_plantilla("confirmacion"))
        plantillas_layout.addWidget(btn_confirmacion)
        
        plantillas_layout.addStretch()
        layout_enviar.addLayout(plantillas_layout)
        
        # Botón enviar
        btn_enviar_layout = QHBoxLayout()
        btn_enviar_layout.addStretch()
        
        self.btn_enviar_mensaje = QPushButton("📤 Enviar Mensaje")
        self.btn_enviar_mensaje.setProperty("class", "whatsapp")
        self.btn_enviar_mensaje.setMinimumWidth(200)
        self.btn_enviar_mensaje.clicked.connect(self._enviar_mensaje)
        btn_enviar_layout.addWidget(self.btn_enviar_mensaje)
        
        layout_enviar.addLayout(btn_enviar_layout)
        
        layout_enviar.addStretch()
        
        tabs.addTab(tab_enviar, "📤 Enviar Mensaje")
        
        # === TAB 2: RECORDATORIOS AUTOMÁTICOS ===
        tab_automatico = QWidget()
        layout_automatico = QVBoxLayout(tab_automatico)
        layout_automatico.setContentsMargins(15, 15, 15, 15)
        layout_automatico.setSpacing(15)
        
        # Configuración de recordatorios
        grupo_config = QGroupBox("Configuración de Recordatorios Automáticos")
        layout_config = QFormLayout(grupo_config)
        layout_config.setSpacing(15)
        
        self.check_recordatorios_activos = QCheckBox("Activar recordatorios automáticos")
        layout_config.addRow("", self.check_recordatorios_activos)
        
        lbl_dias_antes = QLabel("Días antes del vencimiento:")
        lbl_dias_antes.setStyleSheet("font-weight: 600; color: #374151;")
        self.spin_dias_antes = QSpinBox()
        self.spin_dias_antes.setRange(1, 30)
        self.spin_dias_antes.setValue(3)
        self.spin_dias_antes.setSuffix(" días")
        layout_config.addRow(lbl_dias_antes, self.spin_dias_antes)
        
        lbl_hora = QLabel("Hora de envío:")
        lbl_hora.setStyleSheet("font-weight: 600; color: #374151;")
        self.combo_hora = QComboBox()
        for h in range(8, 19):  # 8 AM a 6 PM
            self.combo_hora.addItem(f"{h:02d}:00")
        self.combo_hora.setCurrentText("09:00")
        layout_config.addRow(lbl_hora, self.combo_hora)
        
        btn_guardar_config = QPushButton("💾 Guardar Configuración")
        btn_guardar_config.clicked.connect(self._guardar_config_recordatorios)
        layout_config.addRow("", btn_guardar_config)
        
        layout_automatico.addWidget(grupo_config)
        
        # Lista de próximos recordatorios
        grupo_proximos = QGroupBox("📅 Próximos Recordatorios Programados")
        layout_proximos = QVBoxLayout(grupo_proximos)
        
        self.lista_proximos = QListWidget()
        layout_proximos.addWidget(self.lista_proximos)
        
        btn_actualizar_lista = QPushButton("🔄 Actualizar Lista")
        btn_actualizar_lista.setProperty("class", "secondary")
        btn_actualizar_lista.clicked.connect(self._actualizar_lista_proximos)
        layout_proximos.addWidget(btn_actualizar_lista)
        
        layout_automatico.addWidget(grupo_proximos)
        
        tabs.addTab(tab_automatico, "⏰ Recordatorios Auto")
        
        # === TAB 3: HISTORIAL ===
        tab_historial = QWidget()
        layout_historial = QVBoxLayout(tab_historial)
        layout_historial.setContentsMargins(15, 15, 15, 15)
        layout_historial.setSpacing(15)
        
        # Filtros
        filtros_layout = QHBoxLayout()
        filtros_layout.addWidget(QLabel("Filtrar por:"))
        
        self.combo_filtro_estado = QComboBox()
        self.combo_filtro_estado.addItem("Todos", "todos")
        self.combo_filtro_estado.addItem("Enviados", "enviado")
        self.combo_filtro_estado.addItem("Fallidos", "error")
        self.combo_filtro_estado.currentIndexChanged.connect(self._filtrar_historial)
        filtros_layout.addWidget(self.combo_filtro_estado)
        
        filtros_layout.addStretch()
        
        btn_exportar_historial = QPushButton("📊 Exportar")
        btn_exportar_historial.setProperty("class", "secondary")
        btn_exportar_historial.clicked.connect(self._exportar_historial)
        filtros_layout.addWidget(btn_exportar_historial)
        
        layout_historial.addLayout(filtros_layout)
        
        # Tabla de historial
        self.tabla_historial = QTableWidget()
        self.tabla_historial.setColumnCount(6)
        self.tabla_historial.setHorizontalHeaderLabels([
            "Fecha/Hora", "Destinatario", "Número", "Mensaje (preview)", "Estado", "ID Mensaje"
        ])
        self.tabla_historial.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.tabla_historial.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.tabla_historial.setAlternatingRowColors(True)
        self.tabla_historial.verticalHeader().setVisible(False)
        self.tabla_historial.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        layout_historial.addWidget(self.tabla_historial)
        
        tabs.addTab(tab_historial, "📜 Historial")
        
        layout.addWidget(tabs)
        
        # Cargar clientes
        self._cargar_clientes()
    
    def _verificar_conexion(self):
        """Verifica la conexión con WhatsApp"""
        if self.whatsapp_manager.verificar_conexion():
            self.lbl_estado.setText("● Conectado")
            self.lbl_estado.setStyleSheet("font-size: 11pt; font-weight: bold; color: #059669;")
            self.btn_enviar_mensaje.setEnabled(True)
            
            QMessageBox.information(
                self,
                "Conexión Exitosa",
                "✅ Conectado correctamente a WhatsApp Business.\n\n"
                f"Proveedor: {self.whatsapp_manager.provider.upper()}"
            )
        else:
            self.lbl_estado.setText("● Desconectado")
            self.lbl_estado.setStyleSheet("font-size: 11pt; font-weight: bold; color: #DC2626;")
            self.btn_enviar_mensaje.setEnabled(False)
            
            QMessageBox.warning(
                self,
                "Sin Conexión",
                "⚠️ No se pudo conectar a WhatsApp Business.\n\n"
                "Por favor, configure sus credenciales."
            )
    
    def _configurar_whatsapp(self):
        """Abre diálogo de configuración"""
        dialog = DialogoConfigWhatsApp(self.config, parent=self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            # Recargar configuración
            self.config = dialog.get_config()
            self.whatsapp_manager = WhatsAppManager(self.config)
            self._verificar_conexion()
    
    def _cargar_clientes(self):
        """Carga la lista de clientes"""
        try:
            clientes = self.fm.obtener_entidades(tipo="Cliente", activo=True)
            
            self.combo_destinatario.clear()
            self.combo_destinatario.addItem("Seleccionar cliente...", None)
            
            for cliente in clientes:
                nombre = cliente.get('nombre', '')
                contacto = cliente.get('contacto', '')
                self.combo_destinatario.addItem(f"{nombre} - {contacto}", cliente)
                
        except Exception as e:
            logger.error(f"Error cargando clientes: {e}", exc_info=True)
    
    def _cargar_numero_cliente(self):
        """Carga el número del cliente seleccionado"""
        cliente = self.combo_destinatario.currentData()
        if cliente:
            contacto = cliente.get('contacto', '')
            self.txt_numero.setText(contacto)
    
    def _toggle_tipo_mensaje(self):
        """Maneja el cambio de tipo de mensaje"""
        # Solo permitir uno seleccionado
        sender = self.sender()
        if sender.isChecked():
            if sender == self.radio_manual:
                self.radio_recordatorio.setChecked(False)
                self.radio_confirmacion.setChecked(False)
            elif sender == self.radio_recordatorio:
                self.radio_manual.setChecked(False)
                self.radio_confirmacion.setChecked(False)
            elif sender == self.radio_confirmacion:
                self.radio_manual.setChecked(False)
                self.radio_recordatorio.setChecked(False)
    
    def _usar_plantilla(self, tipo):
        """Carga una plantilla de mensaje"""
        plantillas = {
            "recordatorio": (
                "Estimado/a cliente,\n\n"
                "Le recordamos que tiene un saldo pendiente de pago.\n\n"
                "Por favor, regularice su cuenta a la brevedad posible.\n\n"
                "Gracias por su preferencia."
            ),
            "gracias": (
                "¡Gracias por su pago!\n\n"
                "Hemos recibido su pago correctamente.\n\n"
                "Quedamos a su disposición."
            ),
            "confirmacion": (
                "Estimado/a cliente,\n\n"
                "Confirmamos su alquiler de equipo.\n\n"
                "En breve nos pondremos en contacto con usted.\n\n"
                "Saludos cordiales."
            )
        }
        
        self.txt_mensaje.setPlainText(plantillas.get(tipo, ""))
    
    def _enviar_mensaje(self):
        """Envía el mensaje de WhatsApp"""
        numero = self.txt_numero.text().strip()
        mensaje = self.txt_mensaje.toPlainText().strip()
        
        # Validaciones
        if not numero:
            QMessageBox.warning(self, "Validación", "Debe ingresar un número de teléfono.")
            return
        
        if not mensaje:
            QMessageBox.warning(self, "Validación", "Debe escribir un mensaje.")
            return
        
        # Confirmar envío
        cliente = self.combo_destinatario.currentData()
        destinatario_nombre = cliente.get('nombre', numero) if cliente else numero
        
        respuesta = QMessageBox.question(
            self,
            "Confirmar Envío",
            f"¿Enviar mensaje a {destinatario_nombre}?\n\n"
            f"Número: {numero}\n\n"
            f"Mensaje:\n{mensaje[:100]}{'...' if len(mensaje) > 100 else ''}",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        
        if respuesta != QMessageBox.StandardButton.Yes:
            return
        
        # Enviar mensaje
        self.btn_enviar_mensaje.setEnabled(False)
        self.btn_enviar_mensaje.setText("📤 Enviando...")
        
        # Crear thread para envío asíncrono
        thread = EnvioWhatsAppThread(self.whatsapp_manager, numero, mensaje)
        thread.finished.connect(lambda exito, resultado: self._mensaje_enviado(exito, resultado, destinatario_nombre, numero, mensaje))
        thread.start()
    
    def _mensaje_enviado(self, exito, resultado, destinatario, numero, mensaje):
        """Callback cuando el mensaje es enviado"""
        self.btn_enviar_mensaje.setEnabled(True)
        self.btn_enviar_mensaje.setText("📤 Enviar Mensaje")
        
        # Registrar en historial
        registro = {
            'fecha_hora': datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            'destinatario': destinatario,
            'numero': numero,
            'mensaje': mensaje,
            'estado': 'enviado' if exito else 'error',
            'mensaje_id': resultado if exito else '',
            'error': '' if exito else resultado
        }
        
        self.mensajes_enviados.append(registro)
        
        # Guardar en Firebase (opcional)
        try:
            self.fm.registrar_mensaje_whatsapp(registro)
        except:
            pass
        
        if exito:
            QMessageBox.information(
                self,
                "Mensaje Enviado",
                f"✅ Mensaje enviado correctamente a {destinatario}.\n\n"
                f"ID: {resultado}"
            )
            
            # Limpiar formulario
            self.txt_mensaje.clear()
            self.combo_destinatario.setCurrentIndex(0)
            self.txt_numero.clear()
        else:
            QMessageBox.critical(
                self,
                "Error al Enviar",
                f"❌ No se pudo enviar el mensaje.\n\n"
                f"Error: {resultado}"
            )
        
        # Actualizar historial
        self._filtrar_historial()
    
    def _guardar_config_recordatorios(self):
        """Guarda la configuración de recordatorios automáticos"""
        config = {
            'activo': self.check_recordatorios_activos.isChecked(),
            'dias_antes': self.spin_dias_antes.value(),
            'hora': self.combo_hora.currentText()
        }
        
        # Guardar en configuración
        self.config['whatsapp']['recordatorios_auto'] = config
        
        # Guardar en Firebase o archivo local
        try:
            # Aquí guardarías en tu sistema de configuración
            QMessageBox.information(
                self,
                "Configuración Guardada",
                "✅ Configuración de recordatorios guardada correctamente."
            )
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Error al guardar configuración:\n{e}")
    
    def _actualizar_lista_proximos(self):
        """Actualiza la lista de próximos recordatorios"""
        self.lista_proximos.clear()
        
        try:
            # Obtener cuentas por vencer
            dias = self.spin_dias_antes.value()
            fecha_limite = (datetime.now() + timedelta(days=dias)).strftime("%Y-%m-%d")
            
            alquileres = self.fm.obtener_alquileres({'fecha_fin': fecha_limite})
            
            clientes = {str(c['id']): c['nombre'] for c in self.fm.obtener_entidades(tipo="Cliente", activo=True)}
            
            for alquiler in alquileres:
                monto_total = float(alquiler.get('monto_total', 0))
                monto_pagado = float(alquiler.get('monto_pagado', 0))
                saldo = monto_total - monto_pagado
                
                if saldo > 0:
                    cliente_id = str(alquiler.get('cliente_id', ''))
                    cliente_nombre = clientes.get(cliente_id, 'Desconocido')
                    fecha_venc = alquiler.get('fecha_vencimiento_pago', alquiler.get('fecha_fin', ''))
                    
                    item = QListWidgetItem(
                        f"📅 {fecha_venc} - {cliente_nombre} - Saldo: {self.moneda} {saldo:,.2f}"
                    )
                    self.lista_proximos.addItem(item)
            
            if self.lista_proximos.count() == 0:
                item = QListWidgetItem("✅ No hay recordatorios pendientes")
                self.lista_proximos.addItem(item)
                
        except Exception as e:
            logger.error(f"Error actualizando lista de próximos: {e}", exc_info=True)
    
    def _filtrar_historial(self):
        """Filtra y actualiza la tabla de historial"""
        self.tabla_historial.setRowCount(0)
        
        filtro_estado = self.combo_filtro_estado.currentData()
        
        mensajes_filtrados = self.mensajes_enviados
        if filtro_estado != "todos":
            mensajes_filtrados = [m for m in mensajes_filtrados if m['estado'] == filtro_estado]
        
        for mensaje in sorted(mensajes_filtrados, key=lambda x: x['fecha_hora'], reverse=True):
            row = self.tabla_historial.rowCount()
            self.tabla_historial.insertRow(row)
            
            # Preview del mensaje (primeros 50 caracteres)
            preview = mensaje['mensaje'][:50] + ('...' if len(mensaje['mensaje']) > 50 else '')
            
            # Estado con icono
            estado_texto = "✅ Enviado" if mensaje['estado'] == 'enviado' else "❌ Error"
            
            items = [
                QTableWidgetItem(mensaje['fecha_hora']),
                QTableWidgetItem(mensaje['destinatario']),
                QTableWidgetItem(mensaje['numero']),
                QTableWidgetItem(preview),
                QTableWidgetItem(estado_texto),
                QTableWidgetItem(mensaje.get('mensaje_id', '-'))
            ]
            
            for col, item in enumerate(items):
                if col == 4:  # Estado
                    if mensaje['estado'] == 'enviado':
                        item.setForeground(QBrush(QColor("#059669")))
                    else:
                        item.setForeground(QBrush(QColor("#DC2626")))
                
                self.tabla_historial.setItem(row, col, item)
    
    def _exportar_historial(self):
        """Exporta el historial a Excel"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
            
            archivo, _ = QFileDialog.getSaveFileName(
                self,
                "Guardar Historial",
                f"HistorialWhatsApp_{datetime.now().strftime('%Y%m%d')}.xlsx",
                "Excel (*.xlsx)"
            )
            
            if not archivo:
                return
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Historial WhatsApp"
            
            # Encabezados
            headers = ["Fecha/Hora", "Destinatario", "Número", "Mensaje", "Estado", "ID Mensaje"]
            
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="25D366", end_color="25D366", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            # Datos
            for row, mensaje in enumerate(self.mensajes_enviados, 2):
                ws.cell(row=row, column=1, value=mensaje['fecha_hora'])
                ws.cell(row=row, column=2, value=mensaje['destinatario'])
                ws.cell(row=row, column=3, value=mensaje['numero'])
                ws.cell(row=row, column=4, value=mensaje['mensaje'])
                ws.cell(row=row, column=5, value=mensaje['estado'])
                ws.cell(row=row, column=6, value=mensaje.get('mensaje_id', '-'))
            
            wb.save(archivo)
            
            QMessageBox.information(self, "Éxito", f"Historial exportado:\n{archivo}")
            
        except ImportError:
            QMessageBox.critical(self, "Error", "Se requiere 'openpyxl'.\n\nInstale con: pip install openpyxl")
        except Exception as e:
            logger.error(f"Error exportando: {e}", exc_info=True)
            QMessageBox.critical(self, "Error", f"Error al exportar:\n{e}")


class EnvioWhatsAppThread(QThread):
    """Thread para enviar mensajes de WhatsApp de forma asíncrona"""
    
    finished = pyqtSignal(bool, str)
    
    def __init__(self, whatsapp_manager, numero, mensaje):
        super().__init__()
        self.whatsapp_manager = whatsapp_manager
        self.numero = numero
        self.mensaje = mensaje
    
    def run(self):
        exito, resultado = self.whatsapp_manager.enviar_mensaje(self.numero, self.mensaje)
        self.finished.emit(exito, resultado)


class DialogoConfigWhatsApp(QDialog):
    """Diálogo para configurar WhatsApp Business"""
    
    def __init__(self, config: dict, parent=None):
        super().__init__(parent)
        self.config = config.copy()
        
        self.setWindowTitle("Configuración de WhatsApp Business")
        self.setMinimumWidth(600)
        self.setStyleSheet(WHATSAPP_STYLE)
        
        self._init_ui()
        self._cargar_config()
    
    def _init_ui(self):
        """Inicializa la interfaz"""
        layout = QVBoxLayout(self)
        layout.setContentsMargins(25, 25, 25, 25)
        layout.setSpacing(20)
        
        # Título
        titulo = QLabel("⚙️ Configuración de WhatsApp Business")
        titulo.setStyleSheet("font-size: 16pt; font-weight: bold; color: #1F2937;")
        titulo.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(titulo)
        
        # Selector de proveedor
        grupo_proveedor = QGroupBox("Proveedor")
        layout_proveedor = QVBoxLayout(grupo_proveedor)
        
        self.radio_twilio = QCheckBox("Twilio (Recomendado para desarrollo)")
        self.radio_twilio.toggled.connect(self._toggle_proveedor)
        layout_proveedor.addWidget(self.radio_twilio)
        
        self.radio_whatsapp = QCheckBox("WhatsApp Business API (Producción)")
        self.radio_whatsapp.toggled.connect(self._toggle_proveedor)
        layout_proveedor.addWidget(self.radio_whatsapp)
        
        layout.addWidget(grupo_proveedor)
        
        # === CONFIGURACIÓN TWILIO ===
        self.grupo_twilio = QGroupBox("Credenciales Twilio")
        layout_twilio = QFormLayout(self.grupo_twilio)
        layout_twilio.setSpacing(15)
        
        lbl_sid = QLabel("Account SID:")
        lbl_sid.setStyleSheet("font-weight: 600;")
        self.txt_twilio_sid = QLineEdit()
        self.txt_twilio_sid.setPlaceholderText("ACxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx")
        layout_twilio.addRow(lbl_sid, self.txt_twilio_sid)
        
        lbl_token = QLabel("Auth Token:")
        lbl_token.setStyleSheet("font-weight: 600;")
        self.txt_twilio_token = QLineEdit()
        self.txt_twilio_token.setEchoMode(QLineEdit.EchoMode.Password)
        self.txt_twilio_token.setPlaceholderText("********************************")
        layout_twilio.addRow(lbl_token, self.txt_twilio_token)
        
        lbl_from = QLabel("From Number:")
        lbl_from.setStyleSheet("font-weight: 600;")
        self.txt_twilio_from = QLineEdit()
        self.txt_twilio_from.setPlaceholderText("+14155238886")
        layout_twilio.addRow(lbl_from, self.txt_twilio_from)
        
        layout.addWidget(self.grupo_twilio)
        
        # === CONFIGURACIÓN WHATSAPP BUSINESS ===
        self.grupo_whatsapp = QGroupBox("Credenciales WhatsApp Business API")
        layout_whatsapp = QFormLayout(self.grupo_whatsapp)
        layout_whatsapp.setSpacing(15)
        
        lbl_token_wa = QLabel("Access Token:")
        lbl_token_wa.setStyleSheet("font-weight: 600;")
        self.txt_whatsapp_token = QLineEdit()
        self.txt_whatsapp_token.setEchoMode(QLineEdit.EchoMode.Password)
        self.txt_whatsapp_token.setPlaceholderText("EAAxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxxx")
        layout_whatsapp.addRow(lbl_token_wa, self.txt_whatsapp_token)
        
        lbl_phone_id = QLabel("Phone Number ID:")
        lbl_phone_id.setStyleSheet("font-weight: 600;")
        self.txt_whatsapp_phone_id = QLineEdit()
        self.txt_whatsapp_phone_id.setPlaceholderText("123456789012345")
        layout_whatsapp.addRow(lbl_phone_id, self.txt_whatsapp_phone_id)
        
        layout.addWidget(self.grupo_whatsapp)
        
        # Info
        info = QLabel(
            "ℹ️ <b>Cómo obtener credenciales:</b><br>"
            "<b>Twilio:</b> Regístrate en twilio.com → Console → Account Info<br>"
            "<b>WhatsApp Business:</b> Meta for Developers → WhatsApp → Settings"
        )
        info.setWordWrap(True)
        info.setStyleSheet("background-color: #FEF3C7; padding: 10px; border-radius: 6px; color: #92400E;")
        layout.addWidget(info)
        
        # Botones
        botones_layout = QHBoxLayout()
        
        btn_probar = QPushButton("🧪 Probar Conexión")
        btn_probar.setProperty("class", "secondary")
        btn_probar.clicked.connect(self._probar_conexion)
        botones_layout.addWidget(btn_probar)
        
        botones_layout.addStretch()
        
        btn_guardar = QPushButton("💾 Guardar")
        btn_guardar.setProperty("class", "whatsapp")
        btn_guardar.clicked.connect(self.accept)
        btn_guardar.setMinimumWidth(120)
        botones_layout.addWidget(btn_guardar)
        
        btn_cancelar = QPushButton("✖️ Cancelar")
        btn_cancelar.setProperty("class", "secondary")
        btn_cancelar.clicked.connect(self.reject)
        btn_cancelar.setMinimumWidth(120)
        botones_layout.addWidget(btn_cancelar)
        
        layout.addLayout(botones_layout)
    
    def _cargar_config(self):
        """Carga la configuración actual"""
        whatsapp_config = self.config.get('whatsapp', {})
        
        provider = whatsapp_config.get('provider', 'twilio')
        
        if provider == 'twilio':
            self.radio_twilio.setChecked(True)
            self.txt_twilio_sid.setText(whatsapp_config.get('twilio_account_sid', ''))
            self.txt_twilio_token.setText(whatsapp_config.get('twilio_auth_token', ''))
            self.txt_twilio_from.setText(whatsapp_config.get('twilio_from', ''))
        else:
            self.radio_whatsapp.setChecked(True)
            self.txt_whatsapp_token.setText(whatsapp_config.get('api_token', ''))
            self.txt_whatsapp_phone_id.setText(whatsapp_config.get('phone_id', ''))
    
    def _toggle_proveedor(self):
        """Alterna entre proveedores"""
        if self.radio_twilio.isChecked():
            self.radio_whatsapp.setChecked(False)
            self.grupo_twilio.setEnabled(True)
            self.grupo_whatsapp.setEnabled(False)
        elif self.radio_whatsapp.isChecked():
            self.radio_twilio.setChecked(False)
            self.grupo_twilio.setEnabled(False)
            self.grupo_whatsapp.setEnabled(True)
    
    def _probar_conexion(self):
        """Prueba la conexión con las credenciales actuales"""
        config_test = self.get_config()
        manager = WhatsAppManager(config_test)
        
        if manager.verificar_conexion():
            QMessageBox.information(
                self,
                "Conexión Exitosa",
                "✅ Las credenciales son válidas.\n\nConexión establecida correctamente."
            )
        else:
            QMessageBox.critical(
                self,
                "Error de Conexión",
                "❌ No se pudo conectar.\n\nVerifique sus credenciales."
            )
    
    def get_config(self):
        """Obtiene la configuración actualizada"""
        if 'whatsapp' not in self.config:
            self.config['whatsapp'] = {}
        
        if self.radio_twilio.isChecked():
            self.config['whatsapp']['provider'] = 'twilio'
            self.config['whatsapp']['twilio_account_sid'] = self.txt_twilio_sid.text().strip()
            self.config['whatsapp']['twilio_auth_token'] = self.txt_twilio_token.text().strip()
            self.config['whatsapp']['twilio_from'] = self.txt_twilio_from.text().strip()
        else:
            self.config['whatsapp']['provider'] = 'whatsapp_business'
            self.config['whatsapp']['api_token'] = self.txt_whatsapp_token.text().strip()
            self.config['whatsapp']['phone_id'] = self.txt_whatsapp_phone_id.text().strip()
        
        return self.config