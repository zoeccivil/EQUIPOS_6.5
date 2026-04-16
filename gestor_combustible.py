# gestor_combustible.py

"""
Gestor de Combustible - Control de consumo por equipo.
Lee directamente de la colección 'gastos' filtrando por categoría COMBUSTIBLE.
Cada gasto de combustible tiene el campo 'galones' (unidad estándar en RD).
"""

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QTableWidget, QTableWidgetItem, QHeaderView,
    QComboBox, QMessageBox, QAbstractItemView, QGroupBox, QFrame,
    QFileDialog
)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor
from PyQt6.QtCharts import QChart, QChartView, QLineSeries, QValueAxis, QDateTimeAxis
import logging
from datetime import datetime, timedelta
from firebase_manager import FirebaseManager

logger = logging.getLogger(__name__)

# ─── Estilos ────────────────────────────────────────────────────────────────
COMBUSTIBLE_STYLE = """
QWidget {
    background-color: #F3F4F6;
    font-family: 'Segoe UI';
}
QLabel[class="title"] {
    font-size: 20pt;
    font-weight: bold;
    color: #1F2937;
}
QLabel[class="stat-label"] {
    font-size: 10pt;
    color: #6B7280;
    font-weight: 500;
}
QLabel[class="stat-value"] {
    font-size: 18pt;
    font-weight: bold;
    color: #F59E0B;
}
QFrame[class="stat-card"] {
    background-color: #FFFFFF;
    border: 2px solid #E5E7EB;
    border-radius: 10px;
    padding: 15px;
}
QPushButton {
    background-color: #F59E0B;
    color: white;
    border: none;
    border-radius: 6px;
    padding: 10px 18px;
    font-size: 10pt;
    font-weight: 600;
    min-height: 25px;
}
QPushButton:hover  { background-color: #D97706; }
QPushButton:pressed { background-color: #B45309; }
QPushButton[class="secondary"] {
    background-color: #E5E7EB;
    color: #374151;
}
QPushButton[class="secondary"]:hover { background-color: #D1D5DB; }
QPushButton[class="danger"] {
    background-color: #DC2626;
    color: white;
}
QPushButton[class="danger"]:hover { background-color: #B91C1C; }
QTableWidget {
    background-color: #FFFFFF;
    alternate-background-color: #F9FAFB;
    gridline-color: #E5E7EB;
    selection-background-color: #FEF3C7;
    selection-color: #1F2937;
    border: 1px solid #E5E7EB;
    border-radius: 6px;
}
QTableWidget::item { padding: 8px; }
QHeaderView::section {
    background-color: #1F2937;
    color: #FFFFFF;
    padding: 10px;
    border: none;
    font-weight: 600;
}
QLineEdit, QDoubleSpinBox, QDateEdit, QComboBox {
    background-color: #FFFFFF;
    border: 2px solid #E5E7EB;
    border-radius: 6px;
    padding: 8px 12px;
    color: #1F2937;
    font-size: 10pt;
}
QLineEdit:focus, QDoubleSpinBox:focus, QDateEdit:focus, QComboBox:focus {
    border: 2px solid #F59E0B;
}
QGroupBox {
    background-color: #FFFFFF;
    border: 2px solid #E5E7EB;
    border-radius: 10px;
    margin-top: 14px;
    padding: 15px;
    font-weight: 600;
    font-size: 12pt;
    color: #1F2937;
}
QGroupBox::title {
    subcontrol-origin: margin;
    left: 15px;
    padding: 0 8px;
    background-color: #FFFFFF;
}
"""


class StatCard(QFrame):
    """Tarjeta de estadística reutilizable."""

    def __init__(self, label, value, parent=None):
        super().__init__(parent)
        self.setProperty("class", "stat-card")
        layout = QVBoxLayout(self)
        layout.setSpacing(5)
        lbl = QLabel(label)
        lbl.setProperty("class", "stat-label")
        layout.addWidget(lbl)
        self.lbl_value = QLabel(value)
        self.lbl_value.setProperty("class", "stat-value")
        layout.addWidget(self.lbl_value)

    def actualizar(self, value):
        self.lbl_value.setText(value)


# ─── Widget principal ────────────────────────────────────────────────────────
class GestorCombustible(QWidget):
    """
    Vista de combustible: lee gastos con categoría='COMBUSTIBLE', muestra
    estadísticas en galones y permite registrar nuevas cargas via GastoDialog.
    """

    def __init__(self, fm: FirebaseManager, config: dict, parent=None):
        super().__init__(parent)
        self.fm     = fm
        self.config = config
        self.moneda = config.get('app', {}).get('moneda', 'RD$')

        # Mapas (se populan desde app_gui_qt vía actualizar_mapas o carga propia)
        self.equipos_mapa      = {}
        self.cuentas_mapa      = {}
        self.categorias_mapa   = {}
        self.subcategorias_mapa = {}

        # ID Firestore de la categoría COMBUSTIBLE (se resuelve al cargar datos)
        self._cat_combustible_id = None

        # Gastos cargados (fuente de verdad para los filtros)
        self.gastos = []

        self.setStyleSheet(COMBUSTIBLE_STYLE)
        self._init_ui()

        # Carga inicial en diferido para no bloquear arranque
        from PyQt6.QtCore import QTimer
        QTimer.singleShot(300, self._carga_inicial)

    # ── Interfaz ─────────────────────────────────────────────────────────────
    def _init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(20)

        # Cabecera
        hdr = QHBoxLayout()
        titulo = QLabel("⛽ Control de Combustible")
        titulo.setProperty("class", "title")
        hdr.addWidget(titulo)
        hdr.addStretch()

        hdr.addWidget(QLabel("Equipo:"))
        self.combo_equipo_filtro = QComboBox()
        self.combo_equipo_filtro.addItem("Todos", None)
        self.combo_equipo_filtro.currentIndexChanged.connect(self._aplicar_filtros)
        hdr.addWidget(self.combo_equipo_filtro)

        hdr.addWidget(QLabel("Período:"))
        self.combo_periodo = QComboBox()
        self.combo_periodo.addItem("Último Mes", 30)
        self.combo_periodo.addItem("Últimos 3 Meses", 90)
        self.combo_periodo.addItem("Últimos 6 Meses", 180)
        self.combo_periodo.addItem("Todo", 9999)
        self.combo_periodo.currentIndexChanged.connect(self._cargar_datos)
        hdr.addWidget(self.combo_periodo)

        btn_nueva = QPushButton("➕ Nueva Carga")
        btn_nueva.clicked.connect(self._nueva_carga)
        hdr.addWidget(btn_nueva)

        btn_refresh = QPushButton("🔄 Actualizar")
        btn_refresh.clicked.connect(self._cargar_datos)
        hdr.addWidget(btn_refresh)

        layout.addLayout(hdr)

        # Estadísticas
        stats = QHBoxLayout()
        stats.setSpacing(15)
        self.stat_total_galones = StatCard("Total Galones", "0.000 gal")
        self.stat_total_costo   = StatCard("Costo Total", f"{self.moneda} 0.00")
        self.stat_precio_galon  = StatCard("Precio Promedio/Gal", f"{self.moneda} 0.00")
        self.stat_registros     = StatCard("Registros", "0")
        for card in (self.stat_total_galones, self.stat_total_costo,
                     self.stat_precio_galon, self.stat_registros):
            stats.addWidget(card)
        layout.addLayout(stats)

        # Gráfico
        grupo_grafico = QGroupBox("📈 Historial de Consumo (galones)")
        g_layout = QVBoxLayout(grupo_grafico)
        self.chart_view = self._crear_grafico()
        g_layout.addWidget(self.chart_view)
        layout.addWidget(grupo_grafico)

        # Tabla
        grupo_tabla = QGroupBox("📋 Registro de Cargas")
        t_layout = QVBoxLayout(grupo_tabla)

        acc = QHBoxLayout()
        btn_editar = QPushButton("✏️ Editar")
        btn_editar.clicked.connect(self._editar_carga)
        acc.addWidget(btn_editar)

        btn_eliminar = QPushButton("🗑️ Eliminar")
        btn_eliminar.setProperty("class", "danger")
        btn_eliminar.clicked.connect(self._eliminar_carga)
        acc.addWidget(btn_eliminar)

        btn_exportar = QPushButton("📄 Exportar Excel")
        btn_exportar.setProperty("class", "secondary")
        btn_exportar.clicked.connect(self._exportar_excel)
        acc.addWidget(btn_exportar)
        acc.addStretch()
        t_layout.addLayout(acc)

        self.tabla = QTableWidget()
        self.tabla.setColumnCount(7)
        self.tabla.setHorizontalHeaderLabels([
            "ID", "Fecha", "Equipo", "Galones", "Precio/Gal",
            "Costo Total", "Descripción"
        ])
        self.tabla.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.tabla.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.tabla.setAlternatingRowColors(True)
        self.tabla.verticalHeader().setVisible(False)
        self.tabla.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.tabla.itemDoubleClicked.connect(self._editar_carga)
        t_layout.addWidget(self.tabla)
        layout.addWidget(grupo_tabla)

    def _crear_grafico(self):
        chart = QChart()
        chart.setTitle("Galones por Fecha")
        chart.setAnimationOptions(QChart.AnimationOption.SeriesAnimations)
        chart.setBackgroundBrush(QColor("#FFFFFF"))

        self.series_galones = QLineSeries()
        self.series_galones.setName("Galones")
        chart.addSeries(self.series_galones)

        self.axis_x = QDateTimeAxis()
        self.axis_x.setFormat("dd/MM")
        self.axis_x.setTitleText("Fecha")

        self.axis_y = QValueAxis()
        self.axis_y.setTitleText("Galones")

        chart.addAxis(self.axis_x, Qt.AlignmentFlag.AlignBottom)
        chart.addAxis(self.axis_y, Qt.AlignmentFlag.AlignLeft)
        self.series_galones.attachAxis(self.axis_x)
        self.series_galones.attachAxis(self.axis_y)

        chart.legend().setVisible(True)
        chart.legend().setAlignment(Qt.AlignmentFlag.AlignBottom)

        chart_view = QChartView(chart)
        chart_view.setMinimumHeight(280)
        return chart_view

    # ── Carga de datos ────────────────────────────────────────────────────────
    def _carga_inicial(self):
        """Carga equipos y mapas básicos, luego datos de combustible."""
        self._cargar_mapas_propios()
        self._cargar_datos()

    def _cargar_mapas_propios(self):
        """Carga equipos, categorías, cuentas y subcategorías desde Firebase."""
        try:
            equipos = self.fm.obtener_equipos(activo=None) or []
            self.equipos_mapa = {str(e['id']): e.get('nombre', 'N/A') for e in equipos}
            self.combo_equipo_filtro.clear()
            self.combo_equipo_filtro.addItem("Todos", None)
            for eid, nombre in sorted(self.equipos_mapa.items(), key=lambda x: x[1]):
                self.combo_equipo_filtro.addItem(nombre, eid)
        except Exception as e:
            logger.error(f"GestorCombustible._cargar_mapas_propios equipos: {e}", exc_info=True)

        try:
            self.categorias_mapa = {
                str(k): v for k, v in
                (self.fm.obtener_mapa_global("categorias") or {}).items()
            }
        except Exception as e:
            logger.error(f"GestorCombustible._cargar_mapas_propios categorias: {e}", exc_info=True)

        try:
            self.cuentas_mapa = {
                str(k): v for k, v in
                (self.fm.obtener_mapa_global("cuentas") or {}).items()
            }
        except Exception as e:
            logger.error(f"GestorCombustible._cargar_mapas_propios cuentas: {e}", exc_info=True)

        try:
            self.subcategorias_mapa = {
                str(k): v for k, v in
                (self.fm.obtener_mapa_global("subcategorias") or {}).items()
            }
        except Exception as e:
            logger.error(f"GestorCombustible._cargar_mapas_propios subcategorias: {e}", exc_info=True)

    def actualizar_mapas(self, mapas: dict):
        """
        Llamado por app_gui_qt cuando termina de cargar todos los mapas.
        Refresca la lista de equipos en el filtro y los mapas internos.
        """
        if "equipos" in mapas:
            self.equipos_mapa = mapas["equipos"]
            self.combo_equipo_filtro.blockSignals(True)
            self.combo_equipo_filtro.clear()
            self.combo_equipo_filtro.addItem("Todos", None)
            for eid, nombre in sorted(self.equipos_mapa.items(), key=lambda x: x[1]):
                self.combo_equipo_filtro.addItem(nombre, eid)
            self.combo_equipo_filtro.blockSignals(False)
        if "categorias" in mapas:
            self.categorias_mapa = mapas["categorias"]
        if "cuentas" in mapas:
            self.cuentas_mapa = mapas["cuentas"]
        if "subcategorias" in mapas:
            self.subcategorias_mapa = mapas["subcategorias"]
        # Invalidar caché del id de categoría por si cambió
        self._cat_combustible_id = None

    def _resolver_cat_combustible_id(self) -> str | None:
        """Retorna el Firestore ID de la categoría COMBUSTIBLE (con caché)."""
        if self._cat_combustible_id:
            return self._cat_combustible_id
        # Buscar en el mapa local primero
        for cid, nombre in self.categorias_mapa.items():
            if nombre.upper() == "COMBUSTIBLE":
                self._cat_combustible_id = cid
                return cid
        # Si no está en el mapa, asegurar via Firebase
        try:
            cid = self.fm.ensure_categoria("COMBUSTIBLE")
            if cid:
                self._cat_combustible_id = str(cid)
                return self._cat_combustible_id
        except Exception as e:
            logger.error(f"ensure_categoria COMBUSTIBLE: {e}", exc_info=True)
        return None

    def _cargar_datos(self):
        """Carga gastos de combustible según el período seleccionado."""
        try:
            dias = self.combo_periodo.currentData() or 30
            if dias >= 9999:
                fecha_inicio = "2000-01-01"
            else:
                fecha_inicio = (datetime.now() - timedelta(days=dias)).strftime("%Y-%m-%d")
            fecha_fin = datetime.now().strftime("%Y-%m-%d")

            cat_id = self._resolver_cat_combustible_id()
            if not cat_id:
                logger.warning("No se encontró categoría COMBUSTIBLE")
                self.gastos = []
                self._aplicar_filtros()
                return

            filtros = {
                "fecha_inicio": fecha_inicio,
                "fecha_fin":    fecha_fin,
                "categoria_id": cat_id,
            }
            self.gastos = self.fm.obtener_gastos(filtros)
            self._aplicar_filtros()

        except Exception as e:
            logger.error(f"GestorCombustible._cargar_datos: {e}", exc_info=True)
            QMessageBox.critical(self, "Error", f"Error al cargar datos:\n{e}")

    def _aplicar_filtros(self):
        """Filtra por equipo y actualiza stats, tabla y gráfico."""
        equipo_id = self.combo_equipo_filtro.currentData()
        filtrados = self.gastos
        if equipo_id:
            filtrados = [g for g in filtrados if str(g.get('equipo_id', '')) == equipo_id]

        self._actualizar_estadisticas(filtrados)
        self._actualizar_tabla(filtrados)
        self._actualizar_grafico(filtrados)

    # ── Estadísticas ──────────────────────────────────────────────────────────
    def _actualizar_estadisticas(self, gastos):
        if not gastos:
            self.stat_total_galones.actualizar("0.000 gal")
            self.stat_total_costo.actualizar(f"{self.moneda} 0.00")
            self.stat_precio_galon.actualizar(f"{self.moneda} 0.00")
            self.stat_registros.actualizar("0")
            return

        total_galones = sum(float(g.get('galones') or 0) for g in gastos)
        total_costo   = sum(float(g.get('monto')   or 0) for g in gastos)
        precio_prom   = total_costo / total_galones if total_galones > 0 else 0

        self.stat_total_galones.actualizar(f"{total_galones:,.3f} gal")
        self.stat_total_costo.actualizar(f"{self.moneda} {total_costo:,.2f}")
        self.stat_precio_galon.actualizar(f"{self.moneda} {precio_prom:,.2f}")
        self.stat_registros.actualizar(str(len(gastos)))

    # ── Tabla ─────────────────────────────────────────────────────────────────
    def _actualizar_tabla(self, gastos):
        self.tabla.setRowCount(0)
        for g in sorted(gastos, key=lambda x: x.get('fecha', ''), reverse=True):
            row = self.tabla.rowCount()
            self.tabla.insertRow(row)

            gasto_id   = str(g.get('id', ''))
            fecha      = g.get('fecha', '')
            eq_id      = str(g.get('equipo_id', ''))
            eq_nombre  = self.equipos_mapa.get(eq_id, f"ID:{eq_id}")
            galones    = float(g.get('galones') or 0)
            costo      = float(g.get('monto')   or 0)
            precio_gal = costo / galones if galones > 0 else 0
            descripcion = g.get('descripcion', '') or ''

            self.tabla.setItem(row, 0, QTableWidgetItem(gasto_id))
            self.tabla.setItem(row, 1, QTableWidgetItem(fecha))
            self.tabla.setItem(row, 2, QTableWidgetItem(eq_nombre))

            def right_item(txt):
                it = QTableWidgetItem(txt)
                it.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                return it

            self.tabla.setItem(row, 3, right_item(f"{galones:,.3f}"))
            self.tabla.setItem(row, 4, right_item(f"{self.moneda} {precio_gal:,.2f}"))
            self.tabla.setItem(row, 5, right_item(f"{self.moneda} {costo:,.2f}"))
            self.tabla.setItem(row, 6, QTableWidgetItem(descripcion))

    # ── Gráfico ───────────────────────────────────────────────────────────────
    def _actualizar_grafico(self, gastos):
        self.series_galones.clear()
        if not gastos:
            # Ejes vacíos para que no quede rango extraño
            self.axis_y.setRange(0, 1)
            return

        por_fecha: dict[str, float] = {}
        for g in gastos:
            fecha   = g.get('fecha', '')
            galones = float(g.get('galones') or 0)
            if fecha:
                por_fecha[fecha] = por_fecha.get(fecha, 0) + galones

        if not por_fecha:
            self.axis_y.setRange(0, 1)
            return

        puntos = []
        for fecha, galones in sorted(por_fecha.items()):
            try:
                dt = datetime.strptime(fecha, "%Y-%m-%d")
                ms = int(dt.timestamp() * 1000)
                puntos.append((ms, galones))
                self.series_galones.append(ms, galones)
            except Exception:
                continue

        if not puntos:
            return

        # Actualizar rango del eje X (fechas)
        from PyQt6.QtCore import QDateTime
        min_ms = min(p[0] for p in puntos)
        max_ms = max(p[0] for p in puntos)
        # Si todos los puntos son el mismo día, añadir margen de ±1 día
        if min_ms == max_ms:
            min_ms -= 86_400_000
            max_ms += 86_400_000
        self.axis_x.setRange(
            QDateTime.fromMSecsSinceEpoch(min_ms),
            QDateTime.fromMSecsSinceEpoch(max_ms),
        )

        # Actualizar rango del eje Y (galones), con margen superior del 10 %
        max_gal = max(p[1] for p in puntos)
        self.axis_y.setRange(0, max_gal * 1.10 if max_gal > 0 else 1)

    # ── Acciones ──────────────────────────────────────────────────────────────
    def _nueva_carga(self):
        """Abre GastoDialog pre-seleccionando la categoría COMBUSTIBLE."""
        from dialogos.gasto_dialog import GastoDialog
        cat_id = self._resolver_cat_combustible_id()
        dlg = GastoDialog(
            firebase_manager=self.fm,
            storage_manager=None,
            equipos_mapa=self.equipos_mapa,
            cuentas_mapa=self.cuentas_mapa,
            categorias_mapa=self.categorias_mapa,
            subcategorias_mapa=self.subcategorias_mapa,
            gasto_id=None,
            parent=self,
            moneda_symbol=self.moneda,
        )
        # Pre-seleccionar categoría COMBUSTIBLE
        if cat_id:
            for i in range(dlg.combo_categoria.count()):
                if str(dlg.combo_categoria.itemData(i)) == str(cat_id):
                    dlg.combo_categoria.setCurrentIndex(i)
                    break
        if dlg.exec():
            self._cargar_datos()

    def _editar_carga(self):
        """Abre GastoDialog para editar el gasto seleccionado."""
        current_row = self.tabla.currentRow()
        if current_row < 0:
            QMessageBox.warning(self, "Sin Selección", "Seleccione un registro primero.")
            return

        gasto_id = self.tabla.item(current_row, 0).text()
        from dialogos.gasto_dialog import GastoDialog
        dlg = GastoDialog(
            firebase_manager=self.fm,
            storage_manager=None,
            equipos_mapa=self.equipos_mapa,
            cuentas_mapa=self.cuentas_mapa,
            categorias_mapa=self.categorias_mapa,
            subcategorias_mapa=self.subcategorias_mapa,
            gasto_id=gasto_id,
            parent=self,
            moneda_symbol=self.moneda,
        )
        if dlg.exec():
            self._cargar_datos()

    def _eliminar_carga(self):
        """Elimina el gasto de combustible seleccionado."""
        current_row = self.tabla.currentRow()
        if current_row < 0:
            QMessageBox.warning(self, "Sin Selección", "Seleccione un registro primero.")
            return

        gasto_id = self.tabla.item(current_row, 0).text()
        fecha    = self.tabla.item(current_row, 1).text()
        equipo   = self.tabla.item(current_row, 2).text()

        resp = QMessageBox.question(
            self,
            "Confirmar Eliminación",
            f"¿Eliminar la carga de combustible?\n\n"
            f"Fecha: {fecha}\nEquipo: {equipo}\n\n"
            "Esta acción no se puede deshacer.",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if resp != QMessageBox.StandardButton.Yes:
            return

        try:
            self.fm.eliminar_gasto(gasto_id)
            QMessageBox.information(self, "Éxito", "Registro eliminado correctamente.")
            self._cargar_datos()
        except Exception as e:
            logger.error(f"Error eliminando gasto {gasto_id}: {e}", exc_info=True)
            QMessageBox.critical(self, "Error", f"No se pudo eliminar:\n{e}")

    # ── Exportar Excel ────────────────────────────────────────────────────────
    def _exportar_excel(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            QMessageBox.critical(
                self, "Error",
                "Se requiere 'openpyxl' para exportar.\n\nInstale con: pip install openpyxl"
            )
            return

        archivo, _ = QFileDialog.getSaveFileName(
            self,
            "Guardar Reporte",
            f"Combustible_{datetime.now().strftime('%Y%m%d')}.xlsx",
            "Excel (*.xlsx)"
        )
        if not archivo:
            return

        try:
            equipo_id_filtro = self.combo_equipo_filtro.currentData()
            datos = self.gastos
            if equipo_id_filtro:
                datos = [g for g in datos if str(g.get('equipo_id', '')) == equipo_id_filtro]

            wb = Workbook()
            ws = wb.active
            ws.title = "Combustible"

            headers = ["Fecha", "Equipo", "Galones", "Precio/Gal", "Costo Total", "Descripción"]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            for row, g in enumerate(sorted(datos, key=lambda x: x.get('fecha', '')), 2):
                eq_id    = str(g.get('equipo_id', ''))
                galones  = float(g.get('galones') or 0)
                costo    = float(g.get('monto')   or 0)
                precio   = costo / galones if galones > 0 else 0
                ws.cell(row=row, column=1, value=g.get('fecha', ''))
                ws.cell(row=row, column=2, value=self.equipos_mapa.get(eq_id, f"ID:{eq_id}"))
                ws.cell(row=row, column=3, value=round(galones, 3))
                ws.cell(row=row, column=4, value=round(precio, 2))
                ws.cell(row=row, column=5, value=round(costo, 2))
                ws.cell(row=row, column=6, value=g.get('descripcion', '') or '')

            wb.save(archivo)
            QMessageBox.information(self, "Éxito", f"Reporte exportado:\n{archivo}")

        except Exception as e:
            logger.error(f"Error exportando Excel: {e}", exc_info=True)
            QMessageBox.critical(self, "Error", f"Error al exportar:\n{e}")
